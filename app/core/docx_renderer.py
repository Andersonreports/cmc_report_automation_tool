"""Fill a bundled WES Word template with a ReportData.

Only the user-entered parts are touched: the patient demography table, the
clinical history paragraph and the Sequence data attributes values. All other
content (results, CNV findings, recommendations, methodology, disclaimer,
references, signatures, appendix) is left exactly as in the template.

Two layout-only fixes are applied so the fixed content stays well laid out
when the clinical history is long:
  * the footer's typed "of 4" becomes a live NUMPAGES field;
  * the empty paragraphs padding page 3 are replaced by a page break before
    the appendix, so it always starts on a fresh page.
Elements are located by their text, not their position.
"""
from __future__ import annotations

import copy
import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ..config.templates import get_template
from .models import ReportData
from .os_utils import resource_path

PATIENT_CELLS = [  # (row, col, label, attribute)
    (0, 0, "Patient ID", "patient_id"), (0, 1, "PIN", "pin"),
    (1, 0, "Age", "age"), (1, 1, "Gender", "gender"),
    (2, 0, "Hospital/Clinic", "hospital"), (2, 1, "Sample Number", "sample_number"),
    (3, 0, "Referring Clinician", "referring_clinician"),
    (3, 1, "Sample Collection Date", "collection_date"),
    (4, 0, "Specimen", "specimen"), (4, 1, "Sample Received Date", "received_date"),
    (5, 1, "Report Date", "report_date"),
]


class TemplateError(Exception):
    pass


# ---------------------------------------------------------------------------
# XML helpers
# ---------------------------------------------------------------------------

def _runs(p_el):
    return [r for r in p_el.iter(qn("w:r")) if r.find(qn("w:t")) is not None]


def _text(p_el) -> str:
    return "".join(t.text or "" for t in p_el.iter(qn("w:t")))


def _set_text(p_el, text: str) -> None:
    """Replace a paragraph's text, keeping the first run's formatting.
    Line breaks in `text` become Word line breaks."""
    runs = _runs(p_el)
    if runs:
        first = runs[0]
        for r in runs[1:]:
            r.getparent().remove(r)
        for child in list(first):
            if child.tag != qn("w:rPr"):
                first.remove(child)
    else:
        first = OxmlElement("w:r")
        rpr = p_el.find(qn("w:pPr") + "/" + qn("w:rPr"))
        if rpr is not None:
            first.append(copy.deepcopy(rpr))
        p_el.append(first)
    for i, line in enumerate(text.split("\n")):
        if i:
            first.append(OxmlElement("w:br"))
        t = OxmlElement("w:t")
        t.text = line
        t.set(qn("xml:space"), "preserve")
        first.append(t)


def _find_para(doc, prefix: str, what: str):
    for p in doc.element.body.iter(qn("w:p")):
        if _text(p).strip().lower().startswith(prefix.lower()):
            return p
    raise TemplateError(f"Template is missing: {what}")


def _find_table(doc, first_cell: str, what: str):
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text.strip().lower().startswith(first_cell.lower()):
            return table
    raise TemplateError(f"Template is missing: {what}")


# ---------------------------------------------------------------------------
# User-entered sections
# ---------------------------------------------------------------------------

def _fill_patient(doc, data: ReportData):
    table = _find_table(doc, "Patient ID", "patient details table")
    for row, col, label, attr in PATIENT_CELLS:
        value = (getattr(data.patient, attr, "") or "").strip()
        _set_text(table.rows[row].cells[col].paragraphs[0]._p, f"{label}: {value}")


def _fill_history(doc, data: ReportData):
    head = _find_para(doc, "Clinical History", "'Clinical History' heading")
    body = head.getnext()
    while body is not None and (body.tag != qn("w:p") or not _text(body).strip()):
        body = body.getnext()
    if body is None:
        raise TemplateError("Template is missing the clinical history paragraph")
    _set_text(body, data.clinical_history.strip())


def _fill_qc(doc, data: ReportData):
    table = _find_table(doc, "Total Read", "'Sequence data attributes' table")
    for row, value in ((0, data.total_reads), (1, data.q30)):
        _set_text(table.rows[row].cells[1].paragraphs[0]._p, value.strip())


# ---------------------------------------------------------------------------
# Layout-only fixes
# ---------------------------------------------------------------------------

def _appendix_page_break(doc):
    body = doc.element.body
    try:
        app_head = _find_para(doc, "Appendix 1", "appendix heading")
    except TemplateError:
        return      # template without an appendix
    sig = next((p for p in body if p.tag == qn("w:p")
                and p.find(".//" + qn("a:blip")) is not None), None)
    if sig is None or app_head.getparent() is not body:
        return
    n = sig.getnext()
    while n is not None and n is not app_head:
        nxt = n.getnext()
        if n.tag == qn("w:p") and not _text(n).strip() \
                and n.find(".//" + qn("w:drawing")) is None:
            body.remove(n)
        n = nxt
    ppr = app_head.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        app_head.insert(0, ppr)
    if ppr.find(qn("w:pageBreakBefore")) is not None:
        return
    pb = OxmlElement("w:pageBreakBefore")
    # CT_PPr order: pStyle, keepNext, keepLines, pageBreakBefore, ...
    before = [ppr.find(qn(t)) for t in ("w:keepLines", "w:keepNext", "w:pStyle")]
    before = next((e for e in before if e is not None), None)
    if before is not None:
        before.addnext(pb)
    else:
        ppr.insert(0, pb)


def _field_runs(rpr, instr: str):
    out = []
    for kind in ("begin", "instr", "separate", "text", "end"):
        r = OxmlElement("w:r")
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = "1"
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        r.append(el)
        out.append(r)
    return out


def _fix_page_count(doc):
    """Replace the footer's typed 'of 4' with a live NUMPAGES field."""
    for section in doc.sections:
        for txbx in section.footer._element.iter(qn("w:txbxContent")):
            for p in txbx.iter(qn("w:p")):
                runs = _runs(p)
                for prev, r in zip(runs, runs[1:]):
                    if (prev.find(qn("w:t")).text or "").strip() == "of" and \
                            (r.find(qn("w:t")).text or "").strip().isdigit():
                        for nr in reversed(_field_runs(r.find(qn("w:rPr")),
                                                       " NUMPAGES \\* MERGEFORMAT ")):
                            r.addnext(nr)
                        p.remove(r)
                        break


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def template_path(template_key: str) -> str:
    return resource_path("app", "templates", get_template(template_key).docx_file)


def render_report(data: ReportData):
    path = template_path(data.template_key)
    if not os.path.exists(path):
        raise TemplateError(f"Template file not found: {path}")
    doc = Document(path)
    _fill_patient(doc, data)
    _fill_history(doc, data)
    _fill_qc(doc, data)
    _appendix_page_break(doc)
    _fix_page_count(doc)
    return doc
