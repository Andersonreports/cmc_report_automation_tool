"""Bulk report generation from an Excel sheet.

One report per row, all with the same template, clinical reviewer and report
date. Nothing is saved until the user exports the rows they select. The row's PIN (Anderson ID) fills the patient details from the live
sheets, exactly as Fetch does; the other columns are what the user supplies:

    Anderson ID | Clinical History | Total Read Generated | Data ≥ Q30 | Methodology link

A blank Total Read Generated / Data ≥ Q30 is taken from the sequencing QC
sheet; a blank Methodology link keeps the template's link. The gene coverage
table is the template's.
"""
from __future__ import annotations

import os
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from datetime import datetime

from ..config.templates import get_template
from . import sheet_client
from .models import PatientInfo, ReportData

# Excel column -> field. Headers are matched ignoring case and extra spaces.
COLUMNS = ("Anderson ID", "Clinical History", "Total Read Generated",
           "Data ≥ Q30", "Methodology link")
_ALIASES = {
    "anderson id": "pin", "pin": "pin",
    "clinical history": "history",
    "total read generated": "reads", "total reads": "reads",
    "data ≥ q30": "q30", "data >= q30": "q30", "q30": "q30",
    "methodology link": "url", "gene list link": "url",
}


class BulkError(Exception):
    pass


@dataclass
class BulkRow:
    number: int                    # the row's number in the Excel sheet
    pin: str
    history: str = ""
    reads: str = ""
    q30: str = ""
    url: str = ""


@dataclass
class Prepared:
    """A row with its patient looked up: `data` is the report, or None with
    the reason in `error`. `notes` are things left blank."""
    row: BulkRow
    data: ReportData | None
    error: str = ""
    notes: list[str] = field(default_factory=list)


def _cell_text(cell) -> str:
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, float) and "%" in (cell.number_format or ""):
        return f"{v * 100:.2f} %"
    return " ".join(str(v).split())


def read_rows(path: str) -> list[BulkRow]:
    """The filled rows of the first sheet. Raises BulkError for a file that
    can't be read or has no Anderson ID column."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise BulkError(f"Could not open the Excel file: {e}") from e
    try:
        rows = list(wb.worksheets[0].iter_rows())
    finally:
        wb.close()
    if not rows:
        raise BulkError("The Excel file is empty.")
    header = [_ALIASES.get(re.sub(r"\s+", " ", _cell_text(c).lower()), "") for c in rows[0]]
    if "pin" not in header:
        raise BulkError("The first row must have the headings: " + ", ".join(COLUMNS) + ".")
    out = []
    for i, cells in enumerate(rows[1:], start=2):
        values = {}
        for key, cell in zip(header, cells):
            if key and key not in values:
                # Clinical history keeps its line breaks; the others are one line.
                values[key] = (str(cell.value).strip() if key == "history" and cell.value
                               is not None else _cell_text(cell))
        if not any(values.values()):
            continue                            # blank row
        out.append(BulkRow(i, values.get("pin", "").upper(), values.get("history", ""),
                           _reads(values.get("reads", "")), _q30(values.get("q30", "")),
                           values.get("url", "")))
    if not out:
        raise BulkError("No patient rows found under the headings.")
    return out


def _reads(text: str) -> str:
    """'12.6' -> '12.6 GB'; '9.96GB' -> '9.96 GB'."""
    if re.fullmatch(r"\d+(\.\d+)?", text):
        return f"{text} GB"
    return re.sub(r"(?<=\d)(?=[A-Za-z])", " ", text)


def _q30(text: str) -> str:
    """'96.34' or '96.34%' or '-96.34%' -> '96.34 %'."""
    m = re.search(r"\d+(\.\d+)?", text)
    return f"{m.group(0)} %" if m else text


def build_data(row: BulkRow, template_key: str, reviewer: str, report_date: str):
    """(ReportData, notes) for one row. Raises BulkError when the patient
    can't be looked up - a report without its details would be unusable."""
    if not re.fullmatch(r"[A-Z0-9]{3,6}\d{6,14}", row.pin):
        raise BulkError("Missing or invalid Anderson ID.")
    cfg = get_template(template_key)
    notes = []
    if not sheet_client.patient_configured(cfg):
        raise BulkError("The patient sheet is not set up.")
    try:
        found = sheet_client.fetch_patient(row.pin, cfg)
    except sheet_client.SheetError as e:
        raise BulkError(str(e)) from e
    if found is None:
        raise BulkError(f"{row.pin} was not found in the patient sheet used by the "
                        f"{cfg.name} template.")
    fields = {"hospital": cfg.hospital, "referring_clinician": cfg.referring_clinician,
              "specimen": cfg.specimen}
    fields.update({k: v for k, v in found.items() if v})
    fields["pin"] = row.pin
    if not fields.get("collection_date"):
        fields["collection_date"] = fields.get("received_date", "")
    fields["report_date"] = report_date

    reads, q30 = row.reads, row.q30
    if (not reads or not q30) and sheet_client.qc_configured():
        try:
            qc = sheet_client.fetch_qc(row.pin) or {}
        except sheet_client.SheetError as e:
            qc = {}
            notes.append(f"QC sheet: {e}")
        reads, q30 = reads or qc.get("total_reads", ""), q30 or qc.get("q30", "")
    if not reads:
        notes.append("Total Read Generated is blank")
    if not q30:
        notes.append("Data ≥ Q30 is blank")
    if not row.history:
        notes.append("Clinical History is blank")
    data = ReportData(template_key=template_key,
                      patient=PatientInfo(**{k: fields.get(k, "") for k in
                                             PatientInfo.__dataclass_fields__}),
                      clinical_history=row.history, total_reads=reads, q30=q30,
                      gene_list_url=row.url, reviewer=reviewer)
    return data, notes


def file_name_part(text: str) -> str:
    """'MEL - 12345' -> 'MEL-12345'; 'Baby. Arun K' -> 'Baby_Arun_K'."""
    text = re.sub(r"\s*-\s*", "-", text.strip())
    return re.sub(r"[^A-Za-z0-9-]+", "_", text).strip("_")


def file_stem(data: ReportData) -> str:
    cfg = get_template(data.template_key)
    date = data.patient.report_date or datetime.now().strftime("%d/%m/%Y")
    return "_".join(filter(None, (file_name_part(data.patient.patient_id) or "report",
                                  cfg.file_suffix, date.replace("/", cfg.file_date_sep))))


def prepare(row: BulkRow, template_key: str, reviewer: str, report_date: str) -> Prepared:
    """Look the row's patient up; nothing is written to disk."""
    try:
        data, notes = build_data(row, template_key, reviewer, report_date)
        return Prepared(row, data, "", notes)
    except BulkError as e:
        return Prepared(row, None, str(e))
    except Exception as e:  # noqa: BLE001 - reported per row, the batch goes on
        return Prepared(row, None, f"{type(e).__name__}: {e}")


def stems(items: list[Prepared]) -> list[str]:
    """File names (no extension) for these reports; two rows with the same
    patient and date get the PIN added so they never overwrite each other."""
    used: set[str] = set()
    out = []
    for item in items:
        stem = file_stem(item.data)
        if stem.lower() in used:
            stem = f"{stem}_{item.row.pin}"
        used.add(stem.lower())
        out.append(stem)
    return out


def export(data: ReportData, folder: str, stem: str, kind: str) -> str:
    """Save one report in `folder`. kind: "docx", "pdf" or "both". Returns what
    was saved; raises on failure."""
    from .report_builder import build_report
    docx = os.path.join(folder, stem + ".docx")
    if kind == "pdf":
        tmp = tempfile.mkdtemp(prefix="cmc_bulk_")
        try:
            _, pdf = build_report(data, os.path.join(tmp, stem + ".docx"), tmp)
            shutil.move(pdf, os.path.join(folder, stem + ".pdf"))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
        return stem + ".pdf"
    if kind == "both":
        build_report(data, docx, folder)
        return f"{stem}.docx + .pdf"
    build_report(data, docx)
    return stem + ".docx"
